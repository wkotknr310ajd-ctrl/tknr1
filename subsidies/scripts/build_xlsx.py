"""subsidies/data/*.json から 補助金カレンダー.xlsx を生成する。

使い方:
    python3 subsidies/scripts/build_xlsx.py [--today YYYY-MM-DD]

「状況」「次の締切」列は --today（省略時は日本時間の今日）を基準に計算する。
"""

import argparse
import json
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "subsidies.json"
HISTORY = ROOT / "data" / "history.json"
OUTPUT = ROOT / "補助金カレンダー.xlsx"

MONTHS = [7, 8, 9, 10, 11, 12, 1, 2, 3, 4, 5, 6]
CONFIDENCE = {"確認済", "推定", "要確認"}
JST = timezone(timedelta(hours=9))

FONT = "游ゴシック"
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="center", wrap_text=True)
FILL = {
    "header": PatternFill("solid", fgColor="D9E1F2"),
    "this_month": PatternFill("solid", fgColor="F4B084"),
    "apply": PatternFill("solid", fgColor="BDD7EE"),
    "apply_this_month": PatternFill("solid", fgColor="F8CBAD"),
    "result": PatternFill("solid", fgColor="C6EFCE"),
    "open": PatternFill("solid", fgColor="FFC7CE"),
    "soon": PatternFill("solid", fgColor="FFEB9C"),
    "changed": PatternFill("solid", fgColor="FFF2CC"),
}


def parse_date(s):
    return datetime.strptime(s, "%Y-%m-%d").date() if s else None


def validate(data):
    errors = []
    ids = set()
    for i, s in enumerate(data["subsidies"]):
        where = f"subsidies[{i}] ({s.get('id', '?')})"
        for key in ("id", "category", "name", "calendar", "periods", "confidence", "last_checked"):
            if key not in s:
                errors.append(f"{where}: '{key}' がありません")
        if s.get("id") in ids:
            errors.append(f"{where}: id が重複しています")
        ids.add(s.get("id"))
        if s.get("confidence") not in CONFIDENCE:
            errors.append(f"{where}: confidence は {sorted(CONFIDENCE)} のいずれか")
        for m in s.get("calendar", {}):
            if not m.isdigit() or int(m) not in MONTHS:
                errors.append(f"{where}: calendar のキー '{m}' は 1〜12 の月番号にしてください")
        for p in s.get("periods", []):
            try:
                start, end = parse_date(p.get("start")), parse_date(p.get("end"))
            except ValueError as e:
                errors.append(f"{where}: 日付の形式が不正です ({e})")
                continue
            if not end:
                errors.append(f"{where}: periods の各要素には end が必要です")
            elif start and start > end:
                errors.append(f"{where}: start が end より後になっています")
        try:
            parse_date(s.get("last_checked"))
        except ValueError as e:
            errors.append(f"{where}: last_checked の形式が不正です ({e})")
    if errors:
        sys.exit("subsidies.json に問題があります:\n  " + "\n  ".join(errors))


def fmt_day(d, ref_year=None):
    return f"{d.month}/{d.day}" if d.year == ref_year else f"{d.year}/{d.month}/{d.day}"


def fmt_period(p):
    start, end = parse_date(p.get("start")), parse_date(p["end"])
    span = f"{fmt_day(start)}〜{fmt_day(end, start.year)}" if start else f"〜{fmt_day(end)}"
    text = f"{p['label']}: {span}" if p.get("label") else span
    return text + ("（推定）" if p.get("estimated") else "")


def status_of(s, today):
    """(状況の文言, 次の締切, 塗りつぶしキー) を返す。"""
    periods = [(parse_date(p.get("start")), parse_date(p["end"]), p) for p in s["periods"]]
    upcoming = sorted((end, start, p) for start, end, p in periods if end >= today)
    next_deadline = upcoming[0][0] if upcoming else None
    for end, start, p in upcoming:
        if start is None or start <= today:
            days = (end - today).days
            label = f"募集中（あと{days}日）" if days else "募集中（本日締切）"
            if p.get("estimated"):
                label += "※推定"
            return label, end, "soon" if days <= 14 else "open"
    if upcoming:
        start = min(start for _, start, _ in upcoming)
        return f"募集予定（{fmt_day(start)}〜）", next_deadline, None
    if s["confidence"] == "要確認" or not periods:
        return "要確認", None, None
    return "今期終了・次回未発表", None, None


def build(today):
    data = json.loads(DATA.read_text(encoding="utf-8"))
    history = json.loads(HISTORY.read_text(encoding="utf-8")) if HISTORY.exists() else []
    validate(data)

    fy = data["fiscal_start_year"]
    last_updated = data["last_updated"]
    names = {s["id"]: s["name"] for s in data["subsidies"]}
    changed_now = {}
    for h in history:
        if h["date"] == last_updated:
            changed_now.setdefault(h["id"], []).append(h)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_properties.tabColor = "4472C4"

    ws["B1"] = f"{data['title']}（{fy}年7月〜{fy + 1}年6月）"
    ws["B1"].font = Font(name=FONT, size=14, bold=True)
    ws["P1"] = f"最終更新: {last_updated}　状況は {today.isoformat()} 時点"
    ws["P1"].font = Font(name=FONT, size=10, color="595959")

    headers = ["", "名称・補助限度額"] + [f"{m}月" for m in MONTHS] + [
        "状況", "次の締切", "募集期間（最新）", "確度", "備考", "今回の変更", "公式URL", "最終確認日",
    ]
    for col, text in enumerate(headers, start=2):
        c = ws.cell(row=2, column=col, value=text or None)
        c.font = Font(name=FONT, size=11, bold=True)
        c.alignment = CENTER
        c.border = BORDER
        c.fill = FILL["header"]
    month_col = {m: 4 + i for i, m in enumerate(MONTHS)}
    this_month_col = month_col[today.month] if fy <= today.year <= fy + 1 else None
    if this_month_col:
        ws.cell(row=2, column=this_month_col).fill = FILL["this_month"]

    row = 3
    group_start, group_cat = 3, None
    for s in data["subsidies"]:
        if s["category"] != group_cat:
            if group_cat is not None and row - 1 > group_start:
                ws.merge_cells(start_row=group_start, start_column=2, end_row=row - 1, end_column=2)
            group_start, group_cat = row, s["category"]
            ws.cell(row=row, column=2, value=s["category"])

        ws.cell(row=row, column=3, value=s["name"])
        for m, col in month_col.items():
            mark = s["calendar"].get(str(m))
            c = ws.cell(row=row, column=col, value=mark)
            if mark and mark.startswith("申請"):
                c.fill = FILL["apply_this_month"] if col == this_month_col else FILL["apply"]
            elif mark:
                c.fill = FILL["result"]

        status, deadline, fill = status_of(s, today)
        st = ws.cell(row=row, column=16, value=status)
        if fill:
            st.fill = FILL[fill]
        ws.cell(row=row, column=17, value=deadline)
        ws.cell(row=row, column=18, value="\n".join(fmt_period(p) for p in s["periods"]) or None)
        ws.cell(row=row, column=19, value=s["confidence"])
        ws.cell(row=row, column=20, value=s.get("note"))
        changes = changed_now.get(s["id"], [])
        if changes:
            ch = ws.cell(row=row, column=21, value="\n".join(f"{h['before']} → {h['after']}" for h in changes))
            ch.fill = FILL["changed"]
            ws.cell(row=row, column=3).fill = FILL["changed"]
        url = s.get("official_url") or (s.get("sources") or [None])[0]
        if url:
            u = ws.cell(row=row, column=22, value=url)
            u.hyperlink = url
            u.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        ws.cell(row=row, column=23, value=parse_date(s["last_checked"]))
        row += 1
    if group_cat is not None and row - 1 > group_start:
        ws.merge_cells(start_row=group_start, start_column=2, end_row=row - 1, end_column=2)
    last_row = row - 1

    for r in range(3, last_row + 1):
        for col in range(2, 24):
            c = ws.cell(row=r, column=col)
            c.border = BORDER
            if col != 22:
                c.font = Font(name=FONT, size=11 if col <= 15 else 10)
            c.alignment = CENTER if (col == 2 or 4 <= col <= 17 or col in (19, 23)) else WRAP
        ws.cell(row=r, column=17).number_format = "yyyy/m/d"
        ws.cell(row=r, column=23).number_format = "yyyy/m/d"

    legend = last_row + 2
    ws.cell(row=legend, column=3, value="凡例: 青=申請月　オレンジ=今月　赤=募集中　黄=締切2週間以内／今回の更新で変わった項目").font = Font(name=FONT, size=10)
    ws.cell(row=legend + 1, column=3, value="確度: 確認済=今年度の要項・告知で確認／推定=前年度実績から推定／要確認=今年度の情報が見つからない").font = Font(name=FONT, size=10)

    widths = {"A": 0.5, "B": 7, "C": 62, "P": 18, "Q": 11, "R": 34, "S": 7, "T": 44, "U": 26, "V": 30, "W": 11}
    for m_col in range(4, 16):
        ws.column_dimensions[get_column_letter(m_col)].width = 6.5
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "D3"

    hs = wb.create_sheet("更新履歴")
    h_headers = ["日付", "補助金", "項目", "変更前", "変更後", "理由", "出典"]
    for col, text in enumerate(h_headers, start=1):
        c = hs.cell(row=1, column=col, value=text)
        c.font = Font(name=FONT, bold=True)
        c.fill = FILL["header"]
        c.border = BORDER
        c.alignment = CENTER
    field_names = {"calendar": "申請月", "periods": "募集期間", "name": "名称・限度額", "note": "備考",
                   "official_url": "公式URL", "confidence": "確度", "added": "追加", "removed": "削除"}
    for r, h in enumerate(sorted(history, key=lambda h: h["date"], reverse=True), start=2):
        values = [parse_date(h["date"]), names.get(h["id"], h["id"]), field_names.get(h["field"], h["field"]),
                  h.get("before"), h.get("after"), h.get("reason"), h.get("source")]
        for col, v in enumerate(values, start=1):
            c = hs.cell(row=r, column=col, value=v)
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            c.alignment = WRAP
        hs.cell(row=r, column=1).number_format = "yyyy/m/d"
        if h.get("source"):
            hs.cell(row=r, column=7).hyperlink = h["source"]
    for col, w in zip("ABCDEFG", [11, 40, 10, 22, 22, 50, 40]):
        hs.column_dimensions[col].width = w
    hs.freeze_panes = "A2"

    wb.save(OUTPUT)
    return OUTPUT


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--today", help="状況を計算する基準日 (YYYY-MM-DD)。省略時は日本時間の今日")
    args = parser.parse_args()
    today = parse_date(args.today) if args.today else datetime.now(JST).date()
    out = build(today)
    print(f"generated: {out.relative_to(ROOT.parent)} (as of {today.isoformat()})")


if __name__ == "__main__":
    main()
