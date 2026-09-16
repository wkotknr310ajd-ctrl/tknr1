# -*- coding: utf-8 -*-
"""
既存のExcelシフト表を読み込んで、Webアプリのシフト表に取り込むための処理。

実際の施設のExcelファイルはシートごとにレイアウトがまちまちなので、
「1〜31の連番が横に並んでいる行(日付ヘッダー)」を自動検出したうえで、
氏名の列・職員の行範囲は管理者に確認してもらう2段階の取り込みにしている。
"""
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from security import normalize_name

MIN_DAY_RUN = 20  # これ以上連続した1,2,3...の並びを「日付ヘッダー」とみなす
PREVIEW_ROWS = 40
PREVIEW_COLS = 20


def open_workbook(path):
    return load_workbook(path, data_only=True, read_only=False)


def detect_day_header(ws):
    """1,2,3...の連続する整数が横に並ぶ行を探す。(row, start_col, length) を返す。見つからなければNone。"""
    best = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60)):
        run_start_col = None
        expect = 1
        for cell in row:
            v = cell.value
            if isinstance(v, bool):
                v = None
            if isinstance(v, (int, float)) and int(v) == v and int(v) == expect:
                if run_start_col is None:
                    run_start_col = cell.column
                expect += 1
            else:
                run_len = expect - 1
                if run_len >= MIN_DAY_RUN and (best is None or run_len > best[2]):
                    best = (cell.row, run_start_col, run_len)
                run_start_col = None
                expect = 1
                if isinstance(v, (int, float)) and int(v) == v and int(v) == 1:
                    run_start_col = cell.column
                    expect = 2
        run_len = expect - 1
        if run_start_col is not None and run_len >= MIN_DAY_RUN and (best is None or run_len > best[2]):
            best = (row[0].row, run_start_col, run_len)
    return best


def guess_name_column(header_row, start_col):
    return max(1, start_col - 1)


def guess_staff_rows(ws, header_row, name_col, max_scan=200):
    start = header_row + 1
    end = start
    blanks = 0
    r = start
    found_any = False
    while r <= header_row + max_scan and r <= ws.max_row:
        v = ws.cell(row=r, column=name_col).value
        if v is not None and str(v).strip() != "":
            end = r
            found_any = True
            blanks = 0
        else:
            blanks += 1
            if found_any and blanks >= 3:
                break
        r += 1
    return start, end if found_any else start


def build_preview(ws, max_rows=PREVIEW_ROWS, max_cols=PREVIEW_COLS):
    """先頭の数十行×数十列をHTML表示用の2次元リストにして返す。"""
    n_rows = min(ws.max_row, max_rows)
    n_cols = min(ws.max_column, max_cols)
    rows = []
    for r in range(1, n_rows + 1):
        row_vals = []
        for c in range(1, n_cols + 1):
            v = ws.cell(row=r, column=c).value
            row_vals.append("" if v is None else str(v))
        rows.append({"row": r, "cells": row_vals})
    col_letters = [get_column_letter(c) for c in range(1, n_cols + 1)]
    return rows, col_letters


def column_letter(idx):
    return get_column_letter(idx)


def column_index(letter):
    letter = letter.strip().upper()
    idx = 0
    for ch in letter:
        if not ("A" <= ch <= "Z"):
            raise ValueError("列はアルファベットで指定してください。")
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    if idx <= 0:
        raise ValueError("列を指定してください。")
    return idx


def run_import(conn, logic, path, sheet_name, department_id, header_row, name_col,
                staff_start, staff_end, day_start_col, year, month):
    from security import days_in_month
    from datetime import date as date_cls

    wb = load_workbook(path, data_only=True, read_only=False)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"シート「{sheet_name}」が見つかりません。")
    ws = wb[sheet_name]

    ndays = days_in_month(date_cls(year, month, 1))
    staff_count = 0
    filled_count = 0
    sort_order = conn.execute(
        "SELECT COALESCE(MAX(sort_order), -1) + 1 AS n FROM roster WHERE department_id = ?",
        (department_id,),
    ).fetchone()["n"]

    for r in range(staff_start, staff_end + 1):
        raw_name = ws.cell(row=r, column=name_col).value
        if raw_name is None or str(raw_name).strip() == "":
            continue
        name = normalize_name(str(raw_name))
        if not name:
            continue
        conn.execute(
            "INSERT OR IGNORE INTO roster(department_id, staff_name, sort_order) VALUES (?, ?, ?)",
            (department_id, name, sort_order),
        )
        sort_order += 1
        staff_count += 1
        for d in range(1, ndays + 1):
            col = day_start_col + d - 1
            v = ws.cell(row=r, column=col).value
            if v is None:
                continue
            code = str(v).strip()
            if code == "":
                continue
            target_date_iso = f"{year:04d}-{month:02d}-{d:02d}"
            logic.set_shift_code(conn, department_id, name, target_date_iso, code)
            filled_count += 1
    conn.commit()
    return staff_count, filled_count
